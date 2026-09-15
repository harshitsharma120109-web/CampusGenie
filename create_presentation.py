import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ── 1. CROP SCREENSHOTS PRECISELY ──────────────────────────────────────────
print("[*] Checking and cropping screenshots...")
base_dir = r"C:\Users\manoj\OneDrive\Desktop\CampusGenie"

# Ensure crops exist
if os.path.exists(os.path.join(base_dir, "screenshot_fullview.png")):
    im_full = Image.open(os.path.join(base_dir, "screenshot_fullview.png"))
    crop_dash = im_full.crop((120, 60, 930, 920))
    crop_dash.save(os.path.join(base_dir, "crop_attendance.png"))
    
    crop_tt = im_full.crop((120, 890, 930, 1720))
    crop_tt.save(os.path.join(base_dir, "crop_timetable_opps.png"))

if os.path.exists(os.path.join(base_dir, "screenshot_chat.png")):
    im_chat = Image.open(os.path.join(base_dir, "screenshot_chat.png"))
    crop_c = im_chat.crop((920, 80, 1480, 1020))
    crop_c.save(os.path.join(base_dir, "crop_chat.png"))

if os.path.exists(os.path.join(base_dir, "screenshot_admin.png")):
    im_admin = Image.open(os.path.join(base_dir, "screenshot_admin.png"))
    crop_adm = im_admin.crop((120, 90, 1480, 860))
    crop_adm.save(os.path.join(base_dir, "crop_admin.png"))

print("[+] All image crops ready.")

# ── 2. INITIALIZE PRESENTATION (16:9 Widescreen) ───────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Theme Colors
BG_COLOR = RGBColor(11, 15, 25)         # #0b0f19
CARD_BG = RGBColor(19, 27, 46)          # #131b2e
CARD_BORDER = RGBColor(34, 48, 78)      # #22304e
TEXT_WHITE = RGBColor(248, 250, 252)    # #f8fafc
TEXT_MUTED = RGBColor(148, 163, 184)    # #94a3b8
TEXT_DIM = RGBColor(100, 116, 139)      # #64748b
COLOR_INDIGO = RGBColor(99, 102, 241)   # #6366f1
COLOR_CYAN = RGBColor(6, 182, 212)      # #06b6d4
COLOR_EMERALD = RGBColor(16, 185, 129)  # #10b981
COLOR_ROSE = RGBColor(244, 63, 94)      # #f43f5e
COLOR_AMBER = RGBColor(245, 158, 11)    # #f59e0b

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, tag_text, title_text, tag_color=COLOR_INDIGO):
    tag_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(0.45), Inches(3.6), Inches(0.35)
    )
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(20, 28, 50)
    tag_box.line.color.rgb = tag_color
    tag_box.line.width = Pt(1)
    tf = tag_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = tag_text.upper()
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = tag_color
    p.font.name = 'Segoe UI'
    p.alignment = PP_ALIGN.CENTER

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE
    p_t.font.name = 'Segoe UI'

def add_card_box(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 1: TITLE SLIDE (Grand Hackathon Showcase)
# ══════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s1)

b1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(5.8), Inches(0.42))
b1.fill.solid()
b1.fill.fore_color.rgb = RGBColor(25, 35, 65)
b1.line.color.rgb = COLOR_CYAN
b1.line.width = Pt(1.2)
tf1 = b1.text_frame
tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf1.paragraphs[0]
p.text = "⚡ SKILLUP HACKATHON WITH IBM SKILLSBUILD • STUDENT AI TRACK"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p.font.name = 'Segoe UI'
p.alignment = PP_ALIGN.CENTER

tb_title = s1.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(7.5), Inches(1.0))
tf = tb_title.text_frame
p = tf.paragraphs[0]
p.text = "CampusGenie"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = TEXT_WHITE
p.font.name = 'Segoe UI'

tb_sub = s1.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(7.5), Inches(0.55))
tf = tb_sub.text_frame
p = tf.paragraphs[0]
p.text = "360° AI Student Copilot & Smart Campus ERP Platform"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_INDIGO
p.font.name = 'Segoe UI'

tb_desc = s1.shapes.add_textbox(Inches(1.0), Inches(2.85), Inches(6.8), Inches(1.3))
tf = tb_desc.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = (
    "A unified full-stack campus intelligence portal that puts an end to attendance debarment anxiety, "
    "fragmented college portals, and late-night study hurdles. Featuring a real-time 75% attendance recovery engine, "
    "exam marks grading, live timetable venues, 24/7 CS academic doubt solver, and safe OTC health triage."
)
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.font.name = 'Segoe UI'

add_card_box(s1, Inches(1.0), Inches(4.3), Inches(6.8), Inches(2.5))

info_tb = s1.shapes.add_textbox(Inches(1.2), Inches(4.45), Inches(6.4), Inches(2.2))
tf = info_tb.text_frame
tf.word_wrap = True

def add_info_row(tf, label, val, is_first=False):
    p = tf.paragraphs[0] if is_first else tf.add_paragraph()
    p.space_after = Pt(4)
    run_lbl = p.add_run()
    run_lbl.text = f"• {label}: "
    run_lbl.font.bold = True
    run_lbl.font.size = Pt(11)
    run_lbl.font.color.rgb = COLOR_CYAN
    run_lbl.font.name = 'Segoe UI'
    
    run_val = p.add_run()
    run_val.text = val
    run_val.font.size = Pt(11)
    run_val.font.color.rgb = TEXT_WHITE
    run_val.font.name = 'Segoe UI'

add_info_row(tf, "Developer / Author", "Harshit Sharma (Roll: 22CS1084)", True)
add_info_row(tf, "Branch / Department", "Computer Science & Engineering • 5th Semester")
add_info_row(tf, "Core Technology", "Python 3.12 Flask, Tailwind CSS, Gunicorn, IBM Bob Technology")
add_info_row(tf, "Live Deployed URL", "https://campusgenie-g9pi.onrender.com/")
add_info_row(tf, "GitHub Repository", "https://github.com/harshitsharma120109-web/CampusGenie")

dash_img_path = os.path.join(base_dir, "crop_attendance.png")
if os.path.exists(dash_img_path):
    s1.shapes.add_picture(dash_img_path, Inches(8.1), Inches(1.0), Inches(4.5), Inches(5.8))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 2: PROBLEM STATEMENT & THE CAMPUSGENIE SOLUTION
# ══════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s2)
add_header(s2, "Challenge & Innovation", "The Real Campus Problem & The CampusGenie Solution")

add_card_box(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), border_color=COLOR_ROSE)

tb_prob_title = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(0.5))
p = tb_prob_title.text_frame.paragraphs[0]
p.text = "🚨 The Critical Campus Dilemma"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_ROSE
p.font.name = 'Segoe UI'

tb_prob_content = s2.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.0), Inches(4.2))
tf = tb_prob_content.text_frame
tf.word_wrap = True

problems = [
    ("Strict 75% Attendance Debarment:", "Students face exam debarment without warning. Existing ERPs show delayed percentages with zero predictive math on how to recover from shortages."),
    ("Fragmented & Outdated College Portals:", "Marks are on one notice board, lecture schedules on paper circulars, and attendance on clunky legacy portals that fail on mobile."),
    ("Unanswered Midnight Study Doubts:", "During late-night exam prep, professors are unavailable. Students waste hours browsing scattered forums for quick CS doubt explanations."),
    ("Hostel Late-Night Health Emergencies:", "Campus medical dispensaries close at evening. Hostellers facing sudden fever, acidity, or migraine have no reliable first-aid triage guidance.")
]

for idx, (head, body) in enumerate(problems):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = f"• {head} "
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = TEXT_WHITE
    r1.font.name = 'Segoe UI'
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'

add_card_box(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), border_color=COLOR_EMERALD)

tb_sol_title = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(0.5))
p = tb_sol_title.text_frame.paragraphs[0]
p.text = "💡 The CampusGenie 360° Solution"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD
p.font.name = 'Segoe UI'

tb_sol_content = s2.shapes.add_textbox(Inches(7.1), Inches(2.35), Inches(5.1), Inches(4.2))
tf = tb_sol_content.text_frame
tf.word_wrap = True

solutions = [
    ("Dynamic 75% Mathematical Recovery Engine:", "Automated recovery calculator computes exact consecutive classes needed to return to safety, plus safe skip allowances."),
    ("Unified Glassmorphic 360° Dashboard:", "Attendance, exam marks (Pass/Fail grading), timetable with classroom venues, and curated opportunities in one unified screen."),
    ("24/7 AI Academic Copilot (30+ CS Topics):", "Instant code snippets, time complexities, and intuitive explanations for Data Structures, OS, DBMS, Networks, and OOP."),
    ("Safe Campus Health & Symptom Triage:", "Evidence-backed first-aid OTC guidance (Paracetamol, ORS) and home remedies for 17 campus symptoms with safety guardrails."),
    ("Faculty & Administrative Hub:", "Professors can enroll students, update marks, schedule lectures, and broadcast hackathons with live sync.")
]

for idx, (head, body) in enumerate(solutions):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.space_after = Pt(8)
    r1 = p.add_run()
    r1.text = f"✔ {head} "
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_CYAN
    r1.font.name = 'Segoe UI'
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 3: STUDENT 360° DASHBOARD & MULTI-STUDENT ROSTER
# ══════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s3)
add_header(s3, "Student Experience", "Student 360° Bio & Multi-Student Switcher")

add_card_box(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
tf = tb.text_frame
tf.word_wrap = True

def add_feature_block(tf, icon, title, desc, is_first=False):
    p = tf.paragraphs[0] if is_first else tf.add_paragraph()
    p.space_after = Pt(3)
    r = p.add_run()
    r.text = f"{icon}  {title}\n"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_CYAN
    r.font.name = 'Segoe UI'
    
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'
    p.space_after = Pt(10)

add_feature_block(tf, "👤", "Personalized Student Bio Card", 
                  "Displays student name, unique Roll No (22CS1084), branch (CSE), and semester. Automatically computes avatar monogram initials (HS).", True)
add_feature_block(tf, "🔄", "Seamless Multi-Student Roster Switcher", 
                  "Header dropdown enables one-click switching between active student records (Harshit Sharma, Priya Patel, Jay) for instant demoing and administration.", False)
add_feature_block(tf, "📊", "Dynamic Overall Attendance Badge", 
                  "Real-time percentage calculator. Flags color-coded status pills: '✅ Attendance: 80.0% (Safe)' in emerald green or '⚠️ Shortage' in amber when below 75%.", False)
add_feature_block(tf, "💎", "Responsive Glassmorphism Architecture", 
                  "Built with Tailwind CSS and modern CSS backdrop blur filters. Fully adaptive across mobile smartphones, tablets, and university desktop screens.", False)

if os.path.exists(dash_img_path):
    s3.shapes.add_picture(dash_img_path, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 4: ATTENDANCE SHORTAGE RECOVERY & EXAM MARKS MATRIX
# ══════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s4)
add_header(s4, "Core Innovation", "Subject Attendance Tracking & 75% Recovery Math")

add_card_box(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
tb = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
tf = tb.text_frame
tf.word_wrap = True

add_feature_block(tf, "📐", "The 75% Shortage Recovery Formula", 
                  "When attendance drops below 75%, CampusGenie calculates the exact consecutive lectures required:\n"
                  "Formula: Consecutive Classes Needed = ceil( (0.75 * Total - Attended) / (1 - 0.75) )\n"
                  "Example: In Computer Networks (21/30 = 70%), it alerts: 'Attend 6 consecutive classes to reach 75% target'.", True)

add_feature_block(tf, "🛡️", "Safe Class Skip Allowance", 
                  "For subjects above 75%, it calculates permissible cuts:\n"
                  "Formula: Safe Skips = floor( (Attended - 0.75 * Total) / 0.75 )\n"
                  "Example: 'You can skip up to 3 more classes and stay at 75%'. Prevents unwarranted absenteeism.", False)

add_feature_block(tf, "📝", "Automated Exam Marks & Grading Matrix", 
                  "Tracks mid-term & end-term marks out of 100 with automated Pass/Fail status (40% threshold). In screenshot: OS (78/100 Pass), DSA (85/100 Pass), CN (38/100 Fail).", False)

add_feature_block(tf, "⚡", "Interactive '+ / -' Attendance Simulator", 
                  "Students can tap '+ Present' or '- Absent' to simulate upcoming attendance impacts in real-time before making decisions.", False)

if os.path.exists(dash_img_path):
    s4.shapes.add_picture(dash_img_path, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 5: LIVE TIMETABLE, CLASS VENUES & OPPORTUNITIES BOARD
# ══════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s5)
add_header(s5, "Campus Logistics", "Live Class Timetable, Venues & Opportunities")

add_card_box(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
tf = tb.text_frame
tf.word_wrap = True

add_feature_block(tf, "📍", "Pinpoint Classroom & Lab Venues", 
                  "Solves the common student confusion of running across campus. Clearly lists specific rooms: 'Room 302 (Lecture Block A)', 'Computer Lab 2 (IT Block)', 'Room 205 (Lecture Block B)'.", True)

add_feature_block(tf, "🟢", "Live Ongoing Lecture Pulsing Badge", 
                  "Context-aware status tags: '● Live Now' with an emerald green pulsing animation for ongoing classes, '✅ Completed' for finished lectures, and '⏳ Upcoming' for future slots.", False)

add_feature_block(tf, "📢", "Centralized College Notices Board", 
                  "Official department circulars broadcast directly from faculty. Categorized with color badges: 'Urgent', 'Exam', 'Notice' (e.g. Semester Exam Form deadline 25th September).", False)

add_feature_block(tf, "🏆", "Curated Hackathons & Internship Board", 
                  "Empowers students with career growth. Features verified national competitions (IBM SkillUp Hackathon 2026) with deadlines and direct 'Apply / View' outbound links.", False)

tt_img_path = os.path.join(base_dir, "crop_timetable_opps.png")
if os.path.exists(tt_img_path):
    s5.shapes.add_picture(tt_img_path, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 6: 24/7 AI DUAL COPILOT (ACADEMIC DOUBTS + HEALTH TRIAGE)
# ══════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s6)
add_header(s6, "AI Innovation", "24/7 AI Copilot: Academic Doubt Solver & Health Triage")

add_card_box(s6, Inches(0.8), Inches(1.6), Inches(6.2), Inches(5.2))
tb = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.8), Inches(4.8))
tf = tb.text_frame
tf.word_wrap = True

add_feature_block(tf, "🎓", "30+ Core CS Academic Knowledge Engine", 
                  "Comprehensive doubt solver spanning 8 Computer Science subjects: DSA (QuickSort, MergeSort, Trees, DP, Heaps), Operating Systems (Deadlocks, Paging, Semaphores), DBMS (ACID, Joins, Normalization), Computer Networks (OSI, TCP/UDP), OOP, COA, Automata, and SE. Generates formatted code snippets & complexity metrics.", True)

add_feature_block(tf, "🩺", "Safe Health Triage & OTC First-Aid Engine", 
                  "Specially designed for late-night hostel emergencies. Covers 17 common student health issues (fever, migraine, cold/cough, acidity, eye strain, food poisoning, back pain, stress). Provides immediate first-aid, safe OTC medications (Paracetamol 500mg, ORS, Cetirizine), and non-pharmacological home remedies.", False)

add_feature_block(tf, "⚠️", "Responsible Medical Guardrails", 
                  "Strict safety protocol: Every health recommendation includes explicit disclaimers and emergency escalation advice instructing students to consult the Campus Medical Officer if symptoms exceed 48 hours.", False)

add_feature_block(tf, "⚡", "Instant Interactive Quick Chips", 
                  "One-tap prompt chips for popular questions (Merge Sort, DP, Heap, Greedy, Recursion) enabling instant academic assistance without typing.", False)

chat_img_path = os.path.join(base_dir, "crop_chat.png")
if os.path.exists(chat_img_path):
    s6.shapes.add_picture(chat_img_path, Inches(7.4), Inches(1.6), Inches(5.1), Inches(5.2))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 7: FACULTY ADMINISTRATION HUB & MASTER ROSTER TABLE
# ══════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s7)
add_header(s7, "Administration", "Faculty Management Hub & Live Student Roster")

add_card_box(s7, Inches(0.8), Inches(1.6), Inches(5.4), Inches(5.2))
tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.0), Inches(4.8))
tf = tb.text_frame
tf.word_wrap = True

add_feature_block(tf, "👨‍🏫", "Dedicated Faculty Control Hub", 
                  "A password-free, secure management portal allowing professors, HODs, and academic proctors to manage campus records in real time.", True)

add_feature_block(tf, "🛠️", "6 Rapid Academic Operations", 
                  "1. Enroll Student: Register new student profiles with roll number & branch.\n"
                  "2. Add Subject: Introduce curriculum subjects & assign teachers.\n"
                  "3. Enter Exam Marks: Input test marks with auto Pass/Fail grading.\n"
                  "4. Schedule Classes: Set lecture timings, rooms, and theory/lab type.\n"
                  "5. Post Opportunities: Broadcast verified hackathons & internships.\n"
                  "6. Broadcast Notices: Publish urgent college announcements.", False)

add_feature_block(tf, "📊", "Comprehensive Student × Subject Master Table", 
                  "Provides professors with an executive summary table of all enrolled students (Harshit, Priya, Jay) cross-matched with every subject's attendance %, attended lectures, and exam scores.", False)

add_feature_block(tf, "⚡", "Instant Two-Way Synchronization", 
                  "Any grade, class schedule, or notice updated by faculty reflects instantaneously on student dashboards without requiring database restarts.", False)

admin_img_path = os.path.join(base_dir, "crop_admin.png")
if os.path.exists(admin_img_path):
    s7.shapes.add_picture(admin_img_path, Inches(6.6), Inches(1.6), Inches(5.9), Inches(5.2))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 8: TECHNICAL ARCHITECTURE, IBM BOB IMPACT & FUTURE ROADMAP
# ══════════════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(s8)
add_header(s8, "Engineering & Future", "Technical Architecture, IBM Bob Impact & Roadmap")

col_w = Inches(3.7)
col_h = Inches(5.2)

# Column 1: Tech Stack
add_card_box(s8, Inches(0.8), Inches(1.6), col_w, col_h)
tb1 = s8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.3), Inches(4.8))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "💻 Technical Architecture"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p.font.name = 'Segoe UI'

tech_items = [
    ("Backend Runtime:", "Python 3.12 with lightweight Flask REST microservices."),
    ("Server & Deployment:", "Production WSGI Gunicorn server hosted 24/7 on Render Cloud PaaS."),
    ("Frontend Stack:", "Semantic HTML5, Tailwind CSS, FontAwesome 6, and Vanilla JavaScript (zero bloat, sub-second load times)."),
    ("Data Persistence:", "Atomic JSON document storage with automatic multi-student schema migrations and fallback recovery."),
    ("Security & Integrity:", "CORS protection, input sanitization, and safe evaluation pipelines.")
]

for idx, (head, body) in enumerate(tech_items):
    p = tf1.add_paragraph()
    p.space_after = Pt(6)
    r1 = p.add_run()
    r1.text = f"• {head} "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = TEXT_WHITE
    r1.font.name = 'Segoe UI'
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'

# Column 2: IBM Bob Technology Impact
add_card_box(s8, Inches(4.8), Inches(1.6), col_w, col_h, border_color=COLOR_INDIGO)
tb2 = s8.shapes.add_textbox(Inches(5.0), Inches(1.8), Inches(3.3), Inches(4.8))
tf2 = tb2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "🤖 IBM Bob Technology Role"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_INDIGO
p.font.name = 'Segoe UI'

bob_items = [
    ("Prompt Engineering Acceleration:", "Leveraged IBM Bob to formulate structured prompts for multi-subject CS curriculum parsing and heuristic classification."),
    ("Medical Safety Rule-Sets:", "Utilized IBM Bob to design the triage validation pipeline, ensuring strictly safe OTC medicine dosage constraints (e.g. Paracetamol 500mg, ORS)."),
    ("Data Schema Architecture:", "Designed normalized student-subject-timetable data structures for clean JSON representation."),
    ("Predictive Math Validation:", "Verified formula boundaries for ceiling/floor attendance calculations to eliminate division-by-zero errors.")
]

for idx, (head, body) in enumerate(bob_items):
    p = tf2.add_paragraph()
    p.space_after = Pt(8)
    r1 = p.add_run()
    r1.text = f"✔ {head} "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_CYAN
    r1.font.name = 'Segoe UI'
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'

# Column 3: Future Roadmap
add_card_box(s8, Inches(8.8), Inches(1.6), col_w, col_h)
tb3 = s8.shapes.add_textbox(Inches(9.0), Inches(1.8), Inches(3.3), Inches(4.8))
tf3 = tb3.text_frame
tf3.word_wrap = True

p = tf3.paragraphs[0]
p.text = "🚀 Future Roadmap"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD
p.font.name = 'Segoe UI'

roadmap_items = [
    ("WhatsApp & SMS Automated Alerts:", "Trigger automatic Twilio notifications to students and parents when attendance dips below 75%."),
    ("Voice-Activated AI Assistant:", "Integrate Web Speech API for voice-driven study doubts and venue inquiries in English and Hindi."),
    ("AI Hall Ticket Generator:", "Generate verifiable PDF exam admit cards with cryptographic QR codes only if attendance criteria is satisfied."),
    ("Biometric RFID Hardware Sync:", "Direct hardware integration with classroom ESP32 RFID scanners for automated attendance capture.")
]

for idx, (head, body) in enumerate(roadmap_items):
    p = tf3.add_paragraph()
    p.space_after = Pt(8)
    r1 = p.add_run()
    r1.text = f"★ {head} "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_EMERALD
    r1.font.name = 'Segoe UI'
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = 'Segoe UI'

# ── 3. SAVE PRESENTATION ───────────────────────────────────────────────────
out_path = os.path.join(base_dir, "CampusGenie_Presentation.pptx")
prs.save(out_path)
print(f"[SUCCESS] Presentation generated successfully: {out_path}")
print(f"File size: {os.path.getsize(out_path):,} bytes")
